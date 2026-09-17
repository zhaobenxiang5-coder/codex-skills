from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

from check_decisions import evaluate_decisions, load_reference_ledger, load_supply_ledger  # noqa: E402
from run_readonly_mvp import (  # noqa: E402
    assign_image_groups,
    evaluate_commercial_gate,
    merge_search_results,
    select_review_candidates,
    write_excel,
    write_html,
)
from validate_brief import validate_brief  # noqa: E402


def valid_brief(image_path: Path) -> dict:
    return {
        "version": 1,
        "productDirectionId": "HNS-TEST-001",
        "categoryId": "dog-harness",
        "categoryNameZh": "狗胸背带",
        "targetValueZh": "方便、舒适、好看、安全遛狗",
        "formFingerprint": {
            "nameZh": "马鞍型短背胸背",
            "mustHave": ["宽背片", "双侧快拆"],
            "preferred": ["宽提手"],
            "exclusions": ["康复吊带"],
        },
        "commercialGate": {
            "minCompleteSkuPrice": 15,
            "priceOperator": ">",
            "maxMoq": 1,
            "stockRequired": True,
            "allowedProductForms": ["complete-product", "product-plus-leash"],
            "skuAccessoryOnlyKeywords": ["单独牵引绳", "项圈", "配件"],
            "skuCompleteProductKeywords": ["胸背", "套装"],
        },
        "referenceImages": [str(image_path)],
        "targetReviewCount": 10,
        "detailCandidateLimit": 10,
    }


class MvpTests(unittest.TestCase):
    def test_validate_brief(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / "reference.jpg"
            Image.new("RGB", (80, 60), "white").save(image)
            self.assertEqual(validate_brief(valid_brief(image)), [])
            broken = valid_brief(image)
            broken["commercialGate"]["maxMoq"] = 0
            self.assertTrue(any("maxMoq" in item for item in validate_brief(broken)))

    def test_route_specific_decisions(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            reference = Path(tmp) / "reference.md"
            reference.write_text(
                "# 台账\n\n## 机制指纹台账\n"
                "- accepted-mechanism｜accepted｜已采用\n"
                "- rejected-mechanism｜rejected｜已否决\n\n"
                "## 已明确否决，不再重复\n- Bad Product｜原因\n",
                encoding="utf-8",
            )
            supply = Path(tmp) / "supply.json"
            supply.write_text(json.dumps({"offers": [{"offerId": "111111", "decision": "reject"}]}), encoding="utf-8")
            products, mechanisms = load_reference_ledger(reference)
            offers = load_supply_ledger(supply)
            accepted = evaluate_decisions(
                mode="supply", products=products, mechanisms=mechanisms, supply_offers=offers,
                product_names=[], mechanism_ids=["accepted-mechanism"], offer_ids=[]
            )
            self.assertTrue(accepted["ok"])
            self.assertEqual(accepted["results"][0]["status"], "implementation-target")
            rejected = evaluate_decisions(
                mode="supply", products=products, mechanisms=mechanisms, supply_offers=offers,
                product_names=[], mechanism_ids=["rejected-mechanism"], offer_ids=["111111"]
            )
            self.assertFalse(rejected["ok"])
            self.assertEqual(sum(1 for row in rejected["results"] if row["blocked"]), 2)

    def test_complete_sku_gate_does_not_use_accessory_price(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / "reference.jpg"
            Image.new("RGB", (80, 60), "white").save(image)
            brief = valid_brief(image)
            detail = {
                "minOrderQty": 1,
                "skus": [
                    {"skuId": "rope", "priceRmb": 5, "stock": 100, "attributes": [{"name": "规格", "value": "单独牵引绳"}]},
                    {"skuId": "harness", "priceRmb": 17, "stock": 20, "attributes": [{"name": "颜色", "value": "绿色"}, {"name": "尺码", "value": "L"}]},
                ],
            }
            result = evaluate_commercial_gate(detail, brief)
            self.assertEqual(result["status"], "条件合格")
            self.assertEqual(result["price"], 17)
            self.assertEqual(result["selectedSku"]["skuId"], "harness")
            detail["minOrderQty"] = 2
            self.assertEqual(evaluate_commercial_gate(detail, brief)["status"], "条件不合格")

    def test_offer_dedupe_and_same_image_group(self) -> None:
        rows = [
            {"query": "马鞍胸背", "axis": "版型", "response": {"offers": [
                {"productId": "100001", "title": "A", "supplierName": "甲", "supplierUrl": "https://detail.1688.com/offer/100001.html", "imageUrls": []},
                {"productId": "100002", "title": "B", "supplierName": "乙", "supplierUrl": "https://detail.1688.com/offer/100002.html", "imageUrls": []},
            ]}},
            {"query": "双侧快拆胸背", "axis": "结构", "response": {"offers": [
                {"productId": "100001", "title": "A", "supplierName": "甲", "supplierUrl": "https://detail.1688.com/offer/100001.html", "imageUrls": []},
            ]}},
        ]
        candidates, hits = merge_search_results(rows)
        self.assertEqual(len(candidates), 2)
        self.assertEqual(len(hits), 3)
        self.assertEqual(len(candidates[0]["searchHits"]), 2)
        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / "same.jpg"
            Image.new("RGB", (100, 80), "#19a974").save(image)
            for row in candidates:
                row["localImage"] = str(image)
            assign_image_groups(candidates)
            self.assertEqual(candidates[0]["sameImageGroup"], candidates[1]["sameImageGroup"])

    def test_review_pack_only_uses_detail_checked_offers(self) -> None:
        candidates = [
            {"offerId": "100001", "aiLabel": "基础款", "firstSeen": 1},
            {"offerId": "100002", "aiLabel": "待核验", "firstSeen": 2},
            {"offerId": "100003", "aiLabel": "边界", "firstSeen": 3},
        ]
        details = {"100001": {"ok": True}, "100003": {"ok": False}}
        selected = select_review_candidates(candidates, details, 10)
        self.assertEqual([row["offerId"] for row in selected], ["100003", "100001"])

    def test_excel_and_html_pack(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            evidence_image = root / "技术证据（不用看）" / "产品图" / "100001.jpg"
            evidence_image.parent.mkdir(parents=True)
            Image.new("RGB", (300, 220), "#d9ead3").save(evidence_image)
            candidate = {
                "humanDecision": "",
                "localImage": str(evidence_image),
                "aiLabel": "边界",
                "gate": {"status": "条件合格", "price": 17, "moq": 1, "stock": 20, "selectedSku": {"skuId": "sku-1", "attributes": [{"value": "绿色"}, {"value": "L"}]}, "reasons": []},
                "title": "马鞍式狗胸背",
                "formTypeZh": "马鞍型短背胸背",
                "visibleModules": ["宽背片", "双侧快拆"],
                "borrowablePointZh": "",
                "mechanismFingerprintZh": "普通短背控制不足 → 宽背片承力 → 日常遛狗更稳定",
                "searchHits": [{"axis": "版型", "query": "马鞍胸背", "rank": 1}],
                "supplierName": "测试供应商",
                "location": "浙江省金华市",
                "sameImageGroup": "",
                "offerId": "100001",
                "url": "https://detail.1688.com/offer/100001.html",
                "ledgerNoteZh": "",
                "marketCheckZh": "待复核",
                "capturedAt": "2026-08-21T00:00:00+08:00",
            }
            xlsx = root / "01_差异化现货候选总表.xlsx"
            page = root / "03_差异化现货人工看图总览.html"
            write_excel(xlsx, [candidate])
            write_html(page, [candidate], root)
            book = load_workbook(xlsx)
            sheet = book["差异化现货候选"]
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet["Q2"].value, "100001")
            self.assertEqual(len(sheet.data_validations.dataValidation), 1)
            text = page.read_text(encoding="utf-8")
            self.assertIn('loading="lazy"', text)
            self.assertIn("100001", text)


if __name__ == "__main__":
    unittest.main()
