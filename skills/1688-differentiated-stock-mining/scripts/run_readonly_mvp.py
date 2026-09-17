#!/usr/bin/env python3
"""Read-only 1688 candidate capture and two-file review-pack generator."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from check_decisions import evaluate_decisions, load_reference_ledger, load_supply_ledger
from validate_brief import load_json, validate_brief


ALLOWED_AI_LABELS = {"开发参考", "边界", "基础款", "机制重复", "条件不合格", "待核验"}
LABEL_PRIORITY = {"开发参考": 0, "边界": 1, "待核验": 2, "基础款": 3, "机制重复": 4, "条件不合格": 5}
SAFE_ENDPOINTS = {"/api/supply/1688/search", "/api/supply/1688/detail"}


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def json_dump(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def stable_hash(value: Any) -> str:
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def safe_slug(value: str, *, fallback: str = "item") -> str:
    slug = re.sub(r"[^a-zA-Z0-9\u4e00-\u9fff_-]+", "-", str(value)).strip("-")
    return (slug or fallback)[:80]


def extract_offer_id(value: Any) -> str:
    text = str(value or "")
    match = re.search(r"(?:offer/|productId[=:]?)(\d{6,})", text)
    if match:
        return match.group(1)
    return text.strip() if re.fullmatch(r"\d{6,}", text.strip()) else ""


def unique_strings(values: Iterable[Any]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in seen:
            seen.add(text)
            output.append(text)
    return output


def validate_search_plan(plan: Any) -> list[str]:
    errors: list[str] = []
    if not isinstance(plan, dict) or plan.get("version") != 1:
        return ["搜索计划必须是version=1的JSON对象。"]
    queries = plan.get("queries")
    if not isinstance(queries, list) or not queries:
        return ["搜索计划至少包含一个查询。"]
    if len(queries) > 50:
        errors.append("MVP单次最多执行50个查询。")
    unsupported = {"province", "sort", "ranking", "rankType", "newOnly", "salesOrder"}
    for index, row in enumerate(queries, 1):
        if not isinstance(row, dict):
            errors.append(f"第{index}个查询必须是对象。")
            continue
        query = str(row.get("query", "")).strip()
        axis = str(row.get("axis", "")).strip()
        if len(query) < 2:
            errors.append(f"第{index}个查询至少需要2个字。")
        if not axis:
            errors.append(f"第{index}个查询缺少搜索轴axis。")
        bad = unsupported.intersection(row)
        if bad:
            errors.append(f"第{index}个查询包含当前接口无法证明的字段：{', '.join(sorted(bad))}。")
    return errors


def post_json(base_url: str, endpoint: str, body: dict[str, Any], timeout: int) -> dict[str, Any]:
    if endpoint not in SAFE_ENDPOINTS:
        raise ValueError(f"禁止请求非只读MVP端点：{endpoint}")
    if body.get("apply") is True:
        raise ValueError("只读MVP禁止apply:true")
    request = urllib.request.Request(
        urllib.parse.urljoin(base_url.rstrip("/") + "/", endpoint.lstrip("/")),
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json", "Accept": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = response.read().decode("utf-8")
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")[:500]
        raise RuntimeError(f"HTTP {error.code}: {detail}") from error
    return json.loads(payload)


def cached_post(
    *,
    base_url: str,
    endpoint: str,
    body: dict[str, Any],
    cache_path: Path,
    timeout: int,
    refresh: bool,
    skip_network: bool,
    request_log: list[dict[str, Any]],
) -> dict[str, Any]:
    if cache_path.exists() and not refresh:
        request_log.append({"at": now_iso(), "method": "CACHE", "endpoint": endpoint, "body": body, "cache": str(cache_path)})
        return load_json(cache_path)
    if skip_network:
        raise RuntimeError(f"缓存不存在且已启用--skip-network：{cache_path}")
    started = now_iso()
    result = post_json(base_url, endpoint, body, timeout)
    if result.get("applied") is not False:
        raise RuntimeError(f"端点没有返回明确applied:false：{endpoint}")
    json_dump(cache_path, result)
    request_log.append({"at": started, "method": "POST", "endpoint": endpoint, "body": body, "applied": result.get("applied")})
    return result


def merge_search_results(search_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    candidates: dict[str, dict[str, Any]] = {}
    raw_hits: list[dict[str, Any]] = []
    for row in search_rows:
        query = row["query"]
        axis = row["axis"]
        response = row["response"]
        for rank, offer in enumerate(response.get("offers") or [], 1):
            offer_id = extract_offer_id(offer.get("productId") or offer.get("supplierUrl"))
            if not offer_id:
                continue
            hit = {"query": query, "axis": axis, "rank": rank, "sourceLabel": "1688当前搜索召回"}
            raw_hits.append({"offerId": offer_id, **hit})
            if offer_id not in candidates:
                candidates[offer_id] = {
                    "offerId": offer_id,
                    "url": offer.get("supplierUrl") or f"https://detail.1688.com/offer/{offer_id}.html",
                    "title": str(offer.get("title") or "").strip(),
                    "supplierName": str(offer.get("supplierName") or "").strip(),
                    "summaryPrice": offer.get("purchaseCostRmb") or offer.get("rawPrice"),
                    "summaryMoq": offer.get("minOrderQty"),
                    "imageUrls": unique_strings(offer.get("imageUrls") or []),
                    "supplierAttributes": offer.get("supplierAttributes") or [],
                    "searchHits": [],
                    "firstSeen": len(candidates) + 1,
                }
            candidate = candidates[offer_id]
            candidate["searchHits"].append(hit)
            candidate["imageUrls"] = unique_strings([*candidate.get("imageUrls", []), *(offer.get("imageUrls") or [])])
            if not candidate.get("title"):
                candidate["title"] = str(offer.get("title") or "").strip()
            if not candidate.get("supplierName"):
                candidate["supplierName"] = str(offer.get("supplierName") or "").strip()
    return list(candidates.values()), raw_hits


def attribute_value(attributes: list[dict[str, Any]], names: Iterable[str]) -> str:
    wanted = {str(name).strip() for name in names}
    for row in attributes or []:
        if str(row.get("name", "")).strip() in wanted:
            value = str(row.get("value", "")).strip()
            if value:
                return value
    return ""


def sku_text(sku: dict[str, Any]) -> str:
    chunks: list[str] = []
    for attribute in sku.get("attributes") or []:
        chunks.extend([str(attribute.get("name") or ""), str(attribute.get("value") or "")])
    return " ".join(chunks).strip()


def as_number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        match = re.search(r"\d+(?:\.\d+)?", str(value))
        return float(match.group()) if match else None


def compare_price(value: float, minimum: float, operator: str) -> bool:
    return value > minimum if operator == ">" else value >= minimum


def evaluate_commercial_gate(detail: dict[str, Any] | None, brief: dict[str, Any]) -> dict[str, Any]:
    if not detail:
        return {"status": "待核验", "reasons": ["尚未读取详情"], "selectedSku": None, "price": None, "stock": None, "moq": None}
    gate = brief["commercialGate"]
    minimum = float(gate["minCompleteSkuPrice"])
    operator = gate["priceOperator"]
    max_moq = int(gate["maxMoq"])
    stock_required = bool(gate["stockRequired"])
    accessories = [item.casefold() for item in gate.get("skuAccessoryOnlyKeywords", [])]
    complete_words = [item.casefold() for item in gate.get("skuCompleteProductKeywords", [])]
    moq = detail.get("minOrderQty")
    moq_number = int(moq) if isinstance(moq, (int, float)) and not isinstance(moq, bool) else None
    if moq_number is not None and moq_number > max_moq:
        return {"status": "条件不合格", "reasons": [f"MOQ={moq_number}，超过上限{max_moq}"], "selectedSku": None, "price": None, "stock": None, "moq": moq_number}

    skus = detail.get("skus") or detail.get("skuVariants") or []
    complete_skus: list[dict[str, Any]] = []
    for sku in skus:
        text = sku_text(sku).casefold()
        accessory_hit = any(word and word in text for word in accessories)
        complete_hit = any(word and word in text for word in complete_words)
        if accessory_hit and not complete_hit:
            continue
        complete_skus.append(sku)

    eligible: list[tuple[float, float | None, dict[str, Any]]] = []
    known_prices: list[float] = []
    known_stocks: list[float] = []
    for sku in complete_skus:
        price = as_number(sku.get("priceRmb") if sku.get("priceRmb") is not None else sku.get("price"))
        stock = as_number(sku.get("stock"))
        if price is not None:
            known_prices.append(price)
        if stock is not None:
            known_stocks.append(stock)
        if price is None or not compare_price(price, minimum, operator):
            continue
        if stock_required and (stock is None or stock <= 0):
            continue
        eligible.append((price, stock, sku))

    if eligible and (moq_number is not None or detail.get("minOrderQty") is not None):
        price, stock, sku = sorted(eligible, key=lambda row: (row[0], -(row[1] or 0)))[0]
        return {"status": "条件合格", "reasons": [], "selectedSku": sku, "price": price, "stock": stock, "moq": moq_number}
    if complete_skus and known_prices and all(not compare_price(price, minimum, operator) for price in known_prices):
        return {"status": "条件不合格", "reasons": [f"完整SKU价格均不满足{operator}¥{minimum:g}"], "selectedSku": None, "price": min(known_prices), "stock": max(known_stocks) if known_stocks else None, "moq": moq_number}
    if stock_required and complete_skus and known_stocks and max(known_stocks) <= 0:
        return {"status": "条件不合格", "reasons": ["完整SKU均无库存"], "selectedSku": None, "price": min(known_prices) if known_prices else None, "stock": 0, "moq": moq_number}
    reasons: list[str] = []
    if not complete_skus:
        reasons.append("没有可确认的完整商品SKU")
    if moq_number is None:
        reasons.append("MOQ缺失")
    if complete_skus and not known_prices:
        reasons.append("完整SKU价格缺失")
    if stock_required and complete_skus and not known_stocks:
        reasons.append("库存缺失")
    return {"status": "待核验", "reasons": reasons or ["SKU证据不足"], "selectedSku": None, "price": min(known_prices) if known_prices else None, "stock": max(known_stocks) if known_stocks else None, "moq": moq_number}


def format_sku(sku: dict[str, Any] | None) -> str:
    if not sku:
        return ""
    parts = [str(row.get("value") or "").strip() for row in sku.get("attributes") or []]
    return "｜".join(item for item in parts if item)


def download_image(url: str, path: Path, *, timeout: int = 30) -> bool:
    if path.exists() and path.stat().st_size > 100:
        return True
    if not url:
        return False
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.1688.com/"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            data = response.read()
        with Image.open(BytesIO(data)) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            image.thumbnail((900, 900))
            path.parent.mkdir(parents=True, exist_ok=True)
            image.save(path, "JPEG", quality=90, optimize=True)
        return True
    except Exception:
        return False


def image_hashes(path: Path) -> tuple[str, int | None]:
    if not path.exists():
        return "", None
    with Image.open(path) as image:
        rgb = ImageOps.exif_transpose(image).convert("RGB")
        normalized = rgb.resize((128, 128))
        pixel_hash = hashlib.sha256(normalized.tobytes()).hexdigest()
        gray = rgb.resize((9, 8)).convert("L")
        pixels = list(gray.get_flattened_data() if hasattr(gray, "get_flattened_data") else gray.getdata())
        value = 0
        for row in range(8):
            for col in range(8):
                value = (value << 1) | int(pixels[row * 9 + col] > pixels[row * 9 + col + 1])
        return pixel_hash, value


def hamming(left: int | None, right: int | None) -> int:
    if left is None or right is None:
        return 999
    return (left ^ right).bit_count()


def assign_image_groups(candidates: list[dict[str, Any]]) -> None:
    parent = list(range(len(candidates)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def union(left: int, right: int) -> None:
        left_root, right_root = find(left), find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    for index, row in enumerate(candidates):
        pixel_hash, dhash = image_hashes(Path(row["localImage"])) if row.get("localImage") else ("", None)
        row["pixelHash"] = pixel_hash
        row["dhash"] = dhash
        for other_index in range(index):
            other = candidates[other_index]
            exact = bool(pixel_hash and pixel_hash == other.get("pixelHash"))
            near = bool(dhash is not None and hamming(dhash, other.get("dhash")) <= 1)
            if exact or near:
                union(index, other_index)
    groups: dict[int, list[int]] = {}
    for index in range(len(candidates)):
        groups.setdefault(find(index), []).append(index)
    sequence = 1
    for indexes in groups.values():
        if len(indexes) < 2:
            continue
        label = f"疑似同图组-{sequence:03d}"
        sequence += 1
        for index in indexes:
            candidates[index]["sameImageGroup"] = label


def load_analysis(path: Path | None) -> dict[str, dict[str, Any]]:
    if path is None:
        return {}
    data = load_json(path)
    if not isinstance(data, dict) or data.get("version") != 1 or not isinstance(data.get("items"), list):
        raise ValueError("AI分析文件必须是version=1且包含items数组。")
    output: dict[str, dict[str, Any]] = {}
    for row in data["items"]:
        if not isinstance(row, dict):
            continue
        offer_id = extract_offer_id(row.get("offerId"))
        if not offer_id:
            continue
        label = str(row.get("aiLabel") or "待核验")
        if label not in ALLOWED_AI_LABELS:
            raise ValueError(f"Offer {offer_id} 使用了不允许的AI初判：{label}")
        output[offer_id] = row
    return output


def select_review_candidates(
    candidates: list[dict[str, Any]],
    details: dict[str, dict[str, Any]],
    target_count: int,
) -> list[dict[str, Any]]:
    """只交付本轮已发起详情核验的 Offer，并按判断优先级排序。"""
    verified_pool = [row for row in candidates if row["offerId"] in details]
    ordered = sorted(
        verified_pool,
        key=lambda row: (LABEL_PRIORITY.get(row["aiLabel"], 99), row["firstSeen"]),
    )
    return ordered[:target_count]


def decision_context(
    category_id: str,
    analysis: dict[str, Any],
    offer_id: str,
    reference_ledger: Path | None,
    supply_ledger: Path | None,
) -> tuple[str, str]:
    if category_id != "dog-harness" or reference_ledger is None or not reference_ledger.exists():
        return "", ""
    products, mechanisms = load_reference_ledger(reference_ledger)
    supply = load_supply_ledger(supply_ledger) if supply_ledger else []
    result = evaluate_decisions(
        mode="supply",
        products=products,
        mechanisms=mechanisms,
        supply_offers=supply,
        product_names=[],
        mechanism_ids=[str(analysis.get("mechanismId") or "")] if analysis.get("mechanismId") else [],
        offer_ids=[offer_id],
    )
    notes = [row["actionZh"] for row in result["results"] if row.get("actionZh")]
    offer_decision = next((row for row in result["results"] if row["kind"] == "offer"), None)
    human = {"want": "我要", "pending": "待定", "reject": "不要"}.get((offer_decision or {}).get("status"), "")
    return human, "；".join(unique_strings(notes))


def enrich_candidates(
    candidates: list[dict[str, Any]],
    details: dict[str, dict[str, Any]],
    brief: dict[str, Any],
    analysis_map: dict[str, dict[str, Any]],
    image_dir: Path,
    reference_ledger: Path | None,
    supply_ledger: Path | None,
) -> None:
    for candidate in candidates:
        offer_id = candidate["offerId"]
        detail_response = details.get(offer_id) or {}
        detail = detail_response.get("detail") if detail_response.get("ok") is True else None
        if detail:
            candidate.update({
                "url": detail.get("supplierUrl") or candidate["url"],
                "title": detail.get("title") or candidate["title"],
                "supplierName": detail.get("supplierName") or candidate["supplierName"],
                "imageUrls": unique_strings([*(detail.get("imageUrls") or []), *candidate.get("imageUrls", [])]),
                "supplierAttributes": detail.get("supplierAttributes") or candidate.get("supplierAttributes", []),
                "material": detail.get("material") or "",
                "salesSignal": attribute_value(detail.get("supplierAttributes") or [], ["近一年全网销量", "销量", "成交", "成交量"]),
            })
        candidate["detailVerified"] = bool(detail)
        candidate["detail"] = detail
        candidate["location"] = attribute_value(candidate.get("supplierAttributes") or [], ["发货地", "供应商地区", "所在地区", "产地"])
        candidate["gate"] = evaluate_commercial_gate(detail, brief)
        analysis = analysis_map.get(offer_id, {})
        label = str(analysis.get("aiLabel") or "待核验")
        if candidate["gate"]["status"] == "条件不合格":
            label = "条件不合格"
        human, ledger_note = decision_context(brief["categoryId"], analysis, offer_id, reference_ledger, supply_ledger)
        if human == "不要":
            label = "条件不合格"
        candidate.update({
            "aiLabel": label,
            "formTypeZh": str(analysis.get("formTypeZh") or brief["formFingerprint"]["nameZh"]),
            "visibleModules": unique_strings(analysis.get("visibleModules") or []),
            "mechanismId": str(analysis.get("mechanismId") or ""),
            "mechanismFingerprintZh": str(analysis.get("mechanismFingerprintZh") or ""),
            "borrowablePointZh": str(analysis.get("borrowablePointZh") or ""),
            "marketCheckZh": str(analysis.get("marketCheckZh") or ""),
            "humanDecision": human,
            "ledgerNoteZh": ledger_note,
            "capturedAt": detail_response.get("capturedAt") or now_iso(),
        })
        image_path = image_dir / f"{offer_id}.jpg"
        if download_image((candidate.get("imageUrls") or [""])[0], image_path):
            candidate["localImage"] = str(image_path)
        else:
            candidate["localImage"] = ""
    assign_image_groups(candidates)


def sku_evidence(candidate: dict[str, Any]) -> str:
    gate = candidate["gate"]
    sku = gate.get("selectedSku")
    if sku:
        return f"{format_sku(sku)}｜SKU {sku.get('skuId', '')}".strip("｜")
    return "；".join(gate.get("reasons") or [])


def search_hit_text(candidate: dict[str, Any]) -> str:
    return "；".join(f"{row['axis']}：{row['query']} 当前召回第{row['rank']}位" for row in candidate.get("searchHits") or [])


def write_excel(path: Path, candidates: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "差异化现货候选"
    headers = [
        "人工判断", "产品图", "AI初判", "条件核验", "产品名称", "整款版型", "具体差异模块", "机制指纹",
        "搜索命中", "供应商", "省市", "完整SKU价格", "MOQ", "库存", "SKU证据", "疑似同图厂家组",
        "Offer ID", "1688链接", "台账提示", "市场复核", "抓取时间",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="00A84F")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(*(Side(style="thin", color="B7D7C4") for _ in range(4)))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row_number, item in enumerate(candidates, 2):
        gate = item["gate"]
        price = gate.get("price")
        stock = gate.get("stock")
        values = [
            item.get("humanDecision", ""), "", item["aiLabel"], gate["status"], item.get("title", ""), item.get("formTypeZh", ""),
            "＋".join(item.get("visibleModules") or []) or item.get("borrowablePointZh", ""), item.get("mechanismFingerprintZh", ""),
            search_hit_text(item), item.get("supplierName", ""), item.get("location", ""), f"¥{price:g}" if isinstance(price, (int, float)) else "",
            gate.get("moq") if gate.get("moq") is not None else "", int(stock) if isinstance(stock, (int, float)) else "",
            sku_evidence(item), item.get("sameImageGroup", ""), item["offerId"], item.get("url", ""), item.get("ledgerNoteZh", ""),
            item.get("marketCheckZh", ""), item.get("capturedAt", ""),
        ]
        sheet.append(values)
        sheet.row_dimensions[row_number].height = 92
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
        image_path = item.get("localImage")
        if image_path and Path(image_path).exists():
            image = ExcelImage(image_path)
            ratio = min(150 / max(image.width, 1), 105 / max(image.height, 1))
            image.width = max(1, int(image.width * ratio))
            image.height = max(1, int(image.height * ratio))
            sheet.add_image(image, f"B{row_number}")
        link_cell = sheet.cell(row_number, 18)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"
    widths = [12, 24, 12, 14, 38, 22, 34, 44, 42, 28, 18, 15, 10, 10, 38, 18, 18, 36, 34, 34, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:U{max(sheet.max_row, 2)}"
    sheet.sheet_view.showGridLines = False
    validation = DataValidation(type="list", formula1='"我要,待定,不要"', allow_blank=True)
    validation.error = "请选择：我要、待定或不要"
    validation.errorTitle = "人工判断"
    sheet.add_data_validation(validation)
    validation.add(f"A2:A{max(sheet.max_row, 2)}")
    sheet.conditional_formatting.add(f"A2:A{max(sheet.max_row, 2)}", CellIsRule(operator="equal", formula=['"我要"'], fill=PatternFill("solid", fgColor="D9EAD3")))
    sheet.conditional_formatting.add(f"A2:A{max(sheet.max_row, 2)}", CellIsRule(operator="equal", formula=['"待定"'], fill=PatternFill("solid", fgColor="FFF2CC")))
    sheet.conditional_formatting.add(f"A2:A{max(sheet.max_row, 2)}", CellIsRule(operator="equal", formula=['"不要"'], fill=PatternFill("solid", fgColor="F4CCCC")))
    sheet.auto_filter.ref = f"A1:U{sheet.max_row}"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def html_options(values: Iterable[str]) -> str:
    options = ['<option value="">全部</option>']
    for value in sorted({str(item) for item in values if item}):
        options.append(f'<option value="{html.escape(value)}">{html.escape(value)}</option>')
    return "".join(options)


def write_html(path: Path, candidates: list[dict[str, Any]], output_root: Path) -> None:
    cards: list[str] = []
    for index, item in enumerate(candidates, 1):
        gate = item["gate"]
        relative_image = ""
        if item.get("localImage"):
            relative_image = Path(item["localImage"]).resolve().relative_to(output_root.resolve()).as_posix()
        modules = "＋".join(item.get("visibleModules") or []) or item.get("borrowablePointZh") or "待人工看图提炼"
        price = f"¥{gate['price']:g}" if isinstance(gate.get("price"), (int, float)) else "待核验"
        stock = str(int(gate["stock"])) if isinstance(gate.get("stock"), (int, float)) else "待核验"
        searchable = " ".join([item.get("title", ""), item.get("supplierName", ""), item.get("location", ""), modules, item.get("offerId", "")]).casefold()
        cards.append(f"""
<article class="card" data-search="{html.escape(searchable)}" data-label="{html.escape(item['aiLabel'])}" data-gate="{html.escape(gate['status'])}" data-location="{html.escape(item.get('location',''))}" data-group="{html.escape(item.get('sameImageGroup',''))}">
  <div class="rank">{index:02d}</div>
  <div class="photo">{f'<img loading="lazy" src="{html.escape(relative_image)}" alt="产品图">' if relative_image else '<div class="no-image">图片未缓存</div>'}</div>
  <div class="body">
    <div class="badges"><span>{html.escape(item['aiLabel'])}</span><span>{html.escape(gate['status'])}</span>{f'<span>{html.escape(item.get("sameImageGroup",""))}</span>' if item.get('sameImageGroup') else ''}</div>
    <h2>{html.escape(item.get('title') or '未命名商品')}</h2>
    <p class="modules">{html.escape(modules)}</p>
    <p>{html.escape(item.get('mechanismFingerprintZh') or '机制指纹待人工判断')}</p>
    <dl><dt>供应商</dt><dd>{html.escape(item.get('supplierName') or '待核验')}｜{html.escape(item.get('location') or '地区待核验')}</dd><dt>价格/MOQ/库存</dt><dd>{price}｜MOQ {gate.get('moq') if gate.get('moq') is not None else '待核验'}｜库存 {stock}</dd><dt>搜索命中</dt><dd>{html.escape(search_hit_text(item))}</dd><dt>Offer ID</dt><dd>{html.escape(item['offerId'])}</dd></dl>
    <a href="{html.escape(item.get('url',''))}" target="_blank" rel="noopener noreferrer">打开1688原始商品页</a>
  </div>
</article>""")
    template = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>1688差异化现货人工看图总览</title>
<style>
:root{{--bg:#f4f7f5;--card:#fff;--text:#17231d;--muted:#64726a;--green:#008f4c;--line:#dce7e0}}*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);font-family:-apple-system,BlinkMacSystemFont,"PingFang SC","Microsoft YaHei",sans-serif;color:var(--text)}}header{{position:sticky;top:0;z-index:5;background:rgba(244,247,245,.96);backdrop-filter:blur(12px);padding:22px clamp(18px,4vw,54px);border-bottom:1px solid var(--line)}}h1{{margin:0 0 6px;font-size:clamp(24px,4vw,38px)}}header p{{margin:0;color:var(--muted)}}.filters{{display:grid;grid-template-columns:2fr repeat(4,1fr);gap:10px;margin-top:18px}}input,select{{width:100%;border:1px solid #c9d8cf;border-radius:10px;padding:11px 12px;background:white;color:var(--text)}}main{{padding:26px clamp(18px,4vw,54px) 60px;display:grid;gap:18px}}.card{{position:relative;display:grid;grid-template-columns:minmax(220px,320px) 1fr;gap:22px;background:var(--card);border:1px solid var(--line);border-radius:18px;padding:18px;box-shadow:0 10px 28px rgba(17,67,42,.06)}}.rank{{position:absolute;top:12px;left:12px;background:#10281c;color:white;border-radius:999px;padding:6px 10px;font-weight:700}}.photo{{min-height:240px;display:flex;align-items:center;justify-content:center;background:#f7f9f8;border-radius:13px;overflow:hidden}}.photo img{{width:100%;height:100%;max-height:320px;object-fit:contain}}.no-image{{color:var(--muted)}}.body h2{{margin:10px 0;font-size:21px}}.badges{{display:flex;gap:8px;flex-wrap:wrap}}.badges span{{background:#e6f5ed;color:#087542;padding:5px 9px;border-radius:999px;font-size:13px}}.modules{{font-weight:700;color:#173d2b}}dl{{display:grid;grid-template-columns:110px 1fr;gap:6px 12px}}dt{{color:var(--muted)}}dd{{margin:0}}a{{display:inline-block;margin-top:12px;color:white;background:var(--green);padding:10px 15px;border-radius:10px;text-decoration:none;font-weight:700}}#empty{{display:none;text-align:center;color:var(--muted);padding:60px}}@media(max-width:820px){{.filters{{grid-template-columns:1fr 1fr}}.filters input{{grid-column:1/-1}}.card{{grid-template-columns:1fr}}.photo{{min-height:200px}}dl{{grid-template-columns:90px 1fr}}}}
</style></head><body><header><h1>1688差异化现货人工看图总览</h1><p>共 <strong id="visibleCount">{len(candidates)}</strong> 款；1688只证明供给，最终采用由人工决定。</p><div class="filters"><input id="search" placeholder="搜索产品、供应商、模块或Offer ID"><select id="label">{html_options(item['aiLabel'] for item in candidates)}</select><select id="gate">{html_options(item['gate']['status'] for item in candidates)}</select><select id="location">{html_options(item.get('location','') for item in candidates)}</select><select id="group">{html_options(item.get('sameImageGroup','') for item in candidates)}</select></div></header><main id="cards">{''.join(cards)}<div id="empty">没有符合当前筛选条件的产品。</div></main>
<script>const controls=['search','label','gate','location','group'].map(id=>document.getElementById(id));const cards=[...document.querySelectorAll('.card')];function apply(){{const q=document.getElementById('search').value.trim().toLowerCase();const f={{label:document.getElementById('label').value,gate:document.getElementById('gate').value,location:document.getElementById('location').value,group:document.getElementById('group').value}};let n=0;cards.forEach(c=>{{const ok=(!q||c.dataset.search.includes(q))&&Object.entries(f).every(([k,v])=>!v||c.dataset[k]===v);c.hidden=!ok;if(ok)n++;}});document.getElementById('visibleCount').textContent=n;document.getElementById('empty').style.display=n?'none':'block';}}controls.forEach(c=>c.addEventListener(c.tagName==='INPUT'?'input':'change',apply));</script></body></html>"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(template, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="只读抓取1688候选并生成Excel与看图页")
    parser.add_argument("--brief", type=Path, required=True)
    parser.add_argument("--search-plan", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--analysis", type=Path)
    parser.add_argument("--base-url", default="http://127.0.0.1:4173")
    parser.add_argument("--timeout", type=int, default=90)
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--skip-network", action="store_true")
    parser.add_argument("--reference-ledger", type=Path)
    parser.add_argument("--supply-ledger", type=Path)
    args = parser.parse_args()

    brief = load_json(args.brief)
    errors = validate_brief(brief, check_images=True)
    plan = load_json(args.search_plan)
    errors.extend(validate_search_plan(plan))
    if errors:
        print(json.dumps({"ok": False, "errors": errors}, ensure_ascii=False, indent=2))
        return 2

    skill_root = Path(__file__).resolve().parents[1]
    reference_ledger = args.reference_ledger
    if reference_ledger is None and brief["categoryId"] == "dog-harness":
        reference_ledger = skill_root.parent / "product-reference-mining" / "references" / "dog-harness-ledger.md"
    supply_ledger = args.supply_ledger
    if supply_ledger is None and brief["categoryId"] == "dog-harness":
        supply_ledger = skill_root / "references" / "ledgers" / "dog-harness-supply-ledger.json"

    output = args.output.expanduser().resolve()
    evidence = output / "技术证据（不用看）"
    search_dir = evidence / "搜索原始响应"
    detail_dir = evidence / "详情原始响应"
    image_dir = evidence / "产品图"
    for directory in (search_dir, detail_dir, image_dir):
        directory.mkdir(parents=True, exist_ok=True)

    request_log: list[dict[str, Any]] = []
    search_rows: list[dict[str, Any]] = []
    for index, query_row in enumerate(plan["queries"], 1):
        query = query_row["query"].strip()
        axis = query_row["axis"].strip()
        cache = search_dir / f"{index:02d}_{safe_slug(axis)}_{hashlib.sha1(query.encode()).hexdigest()[:10]}.json"
        body = {"query": query, "limit": 20, "timeoutMs": min(args.timeout * 1000, 90_000)}
        try:
            response = cached_post(base_url=args.base_url, endpoint="/api/supply/1688/search", body=body, cache_path=cache, timeout=args.timeout, refresh=args.refresh, skip_network=args.skip_network, request_log=request_log)
        except Exception as error:
            response = {"ok": False, "applied": False, "localError": str(error), "offers": []}
            json_dump(cache, response)
        search_rows.append({"query": query, "axis": axis, "response": response})

    candidates, raw_hits = merge_search_results(search_rows)
    if not candidates:
        json_dump(evidence / "run.json", {"ok": False, "generatedAt": now_iso(), "reason": "没有召回到有效Offer", "requests": request_log})
        print(json.dumps({"ok": False, "reason": "没有召回到有效Offer", "output": str(output)}, ensure_ascii=False, indent=2))
        return 3

    detail_limit = min(int(brief["detailCandidateLimit"]), len(candidates))
    details: dict[str, dict[str, Any]] = {}
    for candidate in sorted(candidates, key=lambda row: row["firstSeen"])[:detail_limit]:
        offer_id = candidate["offerId"]
        cache = detail_dir / f"{offer_id}.json"
        body = {"productId": offer_id, "apply": False, "timeoutMs": min(args.timeout * 1000, 90_000)}
        try:
            response = cached_post(base_url=args.base_url, endpoint="/api/supply/1688/detail", body=body, cache_path=cache, timeout=args.timeout, refresh=args.refresh, skip_network=args.skip_network, request_log=request_log)
        except Exception as error:
            response = {"ok": False, "applied": False, "localError": str(error)}
            json_dump(cache, response)
        response["capturedAt"] = now_iso()
        details[offer_id] = response

    analysis_map = load_analysis(args.analysis)
    enrich_candidates(candidates, details, brief, analysis_map, image_dir, reference_ledger, supply_ledger)
    # 尚未读取详情的“待核验”商品不能挤掉已完成 SKU 核验和 AI 拆解的商品。
    review = select_review_candidates(candidates, details, int(brief["targetReviewCount"]))
    output.mkdir(parents=True, exist_ok=True)
    excel_path = output / "01_差异化现货候选总表.xlsx"
    html_path = output / "03_差异化现货人工看图总览.html"
    write_excel(excel_path, review)
    write_html(html_path, review, output)

    request_violations = [
        row for row in request_log
        if row.get("method") not in {"POST", "CACHE"}
        or row.get("endpoint") not in SAFE_ENDPOINTS
        or (row.get("body") or {}).get("apply") is True
    ]
    run = {
        "version": 1,
        "ok": not request_violations,
        "mode": "readonly",
        "generatedAt": now_iso(),
        "productDirectionId": brief["productDirectionId"],
        "categoryId": brief["categoryId"],
        "briefPath": str(args.brief.resolve()),
        "briefHash": stable_hash(brief),
        "searchPlanPath": str(args.search_plan.resolve()),
        "searchPlanHash": stable_hash(plan),
        "analysisPath": str(args.analysis.resolve()) if args.analysis else "",
        "counts": {
            "queries": len(plan["queries"]),
            "rawHits": len(raw_hits),
            "uniqueOffers": len(candidates),
            "detailReadbacks": sum(1 for row in details.values() if row.get("ok") is True),
            "reviewRows": len(review),
        },
        "rawHits": raw_hits,
        "candidates": candidates,
        "requests": request_log,
        "writeRequestCount": len(request_violations),
        "requestViolations": request_violations,
        "visibleFiles": [str(excel_path), str(html_path)],
    }
    json_dump(evidence / "run.json", run)
    json_dump(evidence / "request-log.json", request_log)
    result = {
        "ok": run["ok"],
        "output": str(output),
        "excel": str(excel_path),
        "html": str(html_path),
        "counts": run["counts"],
        "writeRequestCount": run["writeRequestCount"],
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if run["ok"] else 4


if __name__ == "__main__":
    sys.exit(main())
